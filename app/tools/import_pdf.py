# app/tools/import_pdf.py
"""PDFApps – TabImport & document conversion engine: convert TXT, Images, Markdown, DOCX, PPTX, XLSX, HTML, EPUB to PDF."""

from __future__ import annotations

import contextlib
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
import importlib.util
from typing import Any

import fitz
from PIL import Image, ImageOps
from PySide6.QtCore import QMarginsF, Qt
from PySide6.QtGui import (
    QPageLayout,
    QPageSize,
    QTextDocument,
)
from PySide6.QtPrintSupport import QPrinter
from PySide6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QListWidget,
    QListWidgetItem,
    QMessageBox,
    QPushButton,
)
import qtawesome as qta

from app.base import BasePage
from app.pdf_io import atomic_pdf_write
from app.i18n import t
from app.utils import section, danger_btn, result_label_style, check_image_size
from app.constants import DESKTOP, TEXT_SEC
from app.widgets import DropFileEdit

_log = logging.getLogger(__name__)


def _has_win32com() -> bool:
    return sys.platform == "win32" and importlib.util.find_spec("win32com") is not None


IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tiff", ".tif", ".ico", ".heic", ".heif", ".svg")
DOC_EXTS = (".docx", ".doc")
PPT_EXTS = (".pptx", ".ppt", ".pptm")
EXCEL_EXTS = (".xlsx", ".xls", ".xlsm")
TEXT_EXTS = (".json", ".txt", ".csv", ".xml", ".yaml", ".yml", ".ini", ".log", ".py", ".js", ".html", ".css", ".md", ".htm")
EBOOK_EXTS = (".epub", ".xps", ".oxps", ".cbz", ".fb2")

_ALL_SUPPORTED_FILTER = (
    "All Files (*);;"
    "All Supported Files (*.txt *.md *.markdown *.docx *.doc *.pptx *.ppt *.xlsx *.xls *.html *.htm *.epub "
    "*.png *.jpg *.jpeg *.bmp *.tiff *.webp *.gif *.pdf);;"
    "Office Documents (*.docx *.doc *.pptx *.ppt *.xlsx *.xls);;"
    "Images (*.png *.jpg *.jpeg *.bmp *.tiff *.webp *.gif *.svg);;"
    "Text & Markdown (*.txt *.md *.markdown *.json *.csv *.xml *.log *.py);;"
    "Web & eBooks (*.html *.htm *.epub *.xps);;"
    "PDF Files (*.pdf)"
)


class _NoContent:
    """Sentinel result: a converter produced a zero-page document."""

    __slots__ = ("skipped",)

    def __init__(self, skipped: int = 0):
        self.skipped = skipped


def _find_onlyoffice_x2t() -> list[str] | None:
    """Locate ONLYOFFICE x2t binary or flatpak command on Linux/macOS/Windows."""
    bin_x2t = shutil.which("x2t")
    if bin_x2t:
        return [bin_x2t]

    candidates = [
        "/opt/onlyoffice/desktopeditors/converter/x2t",
        "/opt/onlyoffice/desktopeditors/x2t",
        "/usr/bin/x2t",
        "/usr/lib/onlyoffice/desktopeditors/converter/x2t",
        "/usr/lib64/onlyoffice/desktopeditors/converter/x2t",
        "/usr/local/bin/x2t",
        "/var/lib/flatpak/app/org.onlyoffice.desktopeditors/current/active/files/bin/x2t",
    ]
    for p in candidates:
        if os.path.isfile(p) and os.access(p, os.X_OK):
            return [p]

    if shutil.which("flatpak"):
        try:
            res = subprocess.run(
                ["flatpak", "info", "org.onlyoffice.desktopeditors"],
                capture_output=True,
                timeout=5,
                check=False,
            )
            if res.returncode == 0:
                return ["flatpak", "run", "--command=x2t", "org.onlyoffice.desktopeditors"]
        except Exception:
            pass

    return None


def _convert_via_onlyoffice(doc_path: str, pdf_path: str) -> bool:
    """Convert Office document to PDF using ONLYOFFICE x2t utility."""
    cmd_base = _find_onlyoffice_x2t()
    if not cmd_base:
        return False
    try:
        abs_doc = os.path.abspath(doc_path)
        abs_pdf = os.path.abspath(pdf_path)
        with tempfile.TemporaryDirectory() as tmp_dir:
            temp_out = os.path.join(tmp_dir, "output.pdf")
            cmd = list(cmd_base) + [abs_doc, temp_out]
            subprocess.run(cmd, capture_output=True, timeout=60, check=False)
            if os.path.isfile(temp_out) and os.path.getsize(temp_out) > 0:
                shutil.copy2(temp_out, abs_pdf)
                return True
    except Exception as exc:
        _log.warning("ONLYOFFICE conversion failed for %s: %s", doc_path, exc)
    return False


def _convert_via_libreoffice(doc_path: str, pdf_path: str) -> bool:
    """Attempt headless conversion using LibreOffice or OpenOffice if available on Linux/macOS/Windows."""
    lo_bin = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo_bin:
        return False
    try:
        abs_doc = os.path.abspath(doc_path)
        abs_pdf = os.path.abspath(pdf_path)
        with tempfile.TemporaryDirectory() as tmp_dir:
            cmd = [
                lo_bin,
                "--headless",
                "--convert-to", "pdf",
                "--outdir", tmp_dir,
                abs_doc,
            ]
            if sys.platform != "win32":
                lo_profile = os.path.join(tmp_dir, "lo_profile")
                cmd.insert(1, f"-env:UserInstallation=file://{lo_profile}")
            subprocess.run(cmd, capture_output=True, timeout=60, check=False)
            stem = os.path.splitext(os.path.basename(abs_doc))[0]
            generated = os.path.join(tmp_dir, f"{stem}.pdf")
            if os.path.isfile(generated) and os.path.getsize(generated) > 0:
                shutil.copy2(generated, abs_pdf)
                return True
            for fname in os.listdir(tmp_dir):
                if fname.lower().endswith(".pdf"):
                    found = os.path.join(tmp_dir, fname)
                    if os.path.isfile(found) and os.path.getsize(found) > 0:
                        shutil.copy2(found, abs_pdf)
                        return True
    except Exception as exc:
        _log.warning("LibreOffice conversion failed for %s: %s", doc_path, exc)
    return False


# ── Office COM & Native Document Converters ──────────────────────────────────

def convert_docx_to_pdf(doc_path: str, pdf_path: str) -> None:
    """Convert DOCX/DOC to PDF using Microsoft Word COM if available, ONLYOFFICE x2t, LibreOffice, or internal python-docx/XML engine."""
    abs_doc = os.path.abspath(doc_path)
    abs_pdf = os.path.abspath(pdf_path)

    if _has_win32com():
        try:
            import pythoncom
            import win32com.client
            try:
                pythoncom.CoInitialize()
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                word.DisplayAlerts = 0  # wdAlertsNone
                try:
                    doc = word.Documents.Open(abs_doc, ReadOnly=True, ConfirmConversions=False)
                    doc.SaveAs2(abs_pdf, FileFormat=17)  # 17 = wdFormatPDF
                    doc.Close(SaveChanges=0)  # 0 = wdDoNotSaveChanges
                    if os.path.isfile(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                        return
                finally:
                    word.Quit()
                    pythoncom.CoUninitialize()
            except Exception as com_err:
                _log.warning("Word COM conversion failed for %s: %s; falling back to internal converter.", doc_path, com_err)
        except ImportError:
            pass

    if _convert_via_onlyoffice(doc_path, pdf_path):
        return

    if _convert_via_libreoffice(doc_path, pdf_path):
        return

    pdf_doc = fitz.open()
    try:
        process_docx(pdf_doc, doc_path)
        if pdf_doc.page_count > 0:
            atomic_pdf_write(pdf_doc, abs_pdf, close_writer=True)
            return
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_doc, "is_closed", False):
                pdf_doc.close()

    raise RuntimeError(f"Could not convert Word document: {doc_path}")


def convert_pptx_to_pdf(ppt_path: str, pdf_path: str) -> None:
    """Convert PPTX/PPT to PDF using Microsoft PowerPoint COM if available, ONLYOFFICE x2t, LibreOffice, or internal python-pptx/XML engine."""
    abs_ppt = os.path.abspath(ppt_path)
    abs_pdf = os.path.abspath(pdf_path)

    if _has_win32com():
        try:
            import pythoncom
            import win32com.client
            try:
                pythoncom.CoInitialize()
                powerpoint = win32com.client.Dispatch("PowerPoint.Application")
                try:
                    try:
                        presentation = powerpoint.Presentations.Open(abs_ppt, ReadOnly=True, WithWindow=False)
                    except Exception:
                        presentation = powerpoint.Presentations.Open(abs_ppt, ReadOnly=True)
                    presentation.SaveAs(abs_pdf, FileFormat=32)  # 32 = ppSaveAsPDF
                    presentation.Close()
                    if os.path.isfile(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                        return
                finally:
                    powerpoint.Quit()
                    pythoncom.CoUninitialize()
            except Exception as com_err:
                _log.warning("PowerPoint COM conversion failed for %s: %s; falling back to internal converter.", ppt_path, com_err)
        except ImportError:
            pass

    if _convert_via_onlyoffice(ppt_path, pdf_path):
        return

    if _convert_via_libreoffice(ppt_path, pdf_path):
        return

    pdf_doc = fitz.open()
    try:
        process_pptx(pdf_doc, ppt_path)
        if pdf_doc.page_count > 0:
            atomic_pdf_write(pdf_doc, abs_pdf, close_writer=True)
            return
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_doc, "is_closed", False):
                pdf_doc.close()

    raise RuntimeError(f"Could not convert PowerPoint presentation: {ppt_path}")


def convert_xlsx_to_pdf(xlsx_path: str, pdf_path: str) -> None:
    """Convert XLSX/XLS to PDF using Microsoft Excel COM if available, ONLYOFFICE x2t, LibreOffice, or internal openpyxl/XML engine."""
    abs_xlsx = os.path.abspath(xlsx_path)
    abs_pdf = os.path.abspath(pdf_path)

    if _has_win32com():
        try:
            import pythoncom
            import win32com.client
            try:
                pythoncom.CoInitialize()
                excel = win32com.client.Dispatch("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                try:
                    wb = excel.Workbooks.Open(abs_xlsx, ReadOnly=True, UpdateLinks=0)
                    wb.ExportAsFixedFormat(0, abs_pdf)  # 0 = xlTypePDF
                    wb.Close(SaveChanges=False)
                    if os.path.isfile(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                        return
                finally:
                    excel.Quit()
                    pythoncom.CoUninitialize()
            except Exception as com_err:
                _log.warning("Excel COM conversion failed for %s: %s; falling back to internal converter.", xlsx_path, com_err)
        except ImportError:
            pass

    if _convert_via_onlyoffice(xlsx_path, pdf_path):
        return

    if _convert_via_libreoffice(xlsx_path, pdf_path):
        return

    pdf_doc = fitz.open()
    try:
        process_xlsx(pdf_doc, xlsx_path)
        if pdf_doc.page_count > 0:
            atomic_pdf_write(pdf_doc, abs_pdf, close_writer=True)
            return
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_doc, "is_closed", False):
                pdf_doc.close()

    raise RuntimeError(f"Could not convert Excel spreadsheet: {xlsx_path}")


def convert_image_to_pdf(image_path: str, pdf_path: str) -> None:
    """Convert raster and vector images to PDF with EXIF rotation and transparency handling."""
    abs_img = os.path.abspath(image_path)
    abs_pdf = os.path.abspath(pdf_path)

    try:
        with Image.open(abs_img) as img:
            img = ImageOps.exif_transpose(img)
            if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
                background = Image.new("RGB", img.size, (255, 255, 255))
                if img.mode == "RGBA":
                    background.paste(img, mask=img.split()[3])
                elif img.mode == "LA":
                    background.paste(img, mask=img.split()[1])
                else:
                    rgba = img.convert("RGBA")
                    background.paste(rgba, mask=rgba.split()[3])
                img = background
            else:
                img = img.convert("RGB")
            img.save(abs_pdf, "PDF", resolution=150.0)
            if os.path.isfile(abs_pdf) and os.path.getsize(abs_pdf) > 0:
                return
    except Exception as pil_err:
        _log.warning("Pillow image conversion failed for %s: %s; falling back to PyMuPDF.", image_path, pil_err)

    pdf_doc = fitz.open()
    try:
        if process_image(pdf_doc, image_path):
            atomic_pdf_write(pdf_doc, abs_pdf, close_writer=True)
            return
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_doc, "is_closed", False):
                pdf_doc.close()

    raise RuntimeError(f"Could not convert image: {image_path}")


def convert_text_to_pdf(text_path: str, pdf_path: str) -> None:
    """Convert text, JSON, HTML, Markdown, and code to PDF using QTextDocument and QPrinter."""
    with open(text_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    ext = os.path.splitext(text_path)[1].lower()
    doc = QTextDocument()

    if ext == ".json":
        is_html_json = False
        html_content = ""
        try:
            parsed = json.loads(content)
            if isinstance(parsed, dict) and parsed.get("metadata", {}).get("format") == "html" and "data" in parsed:
                is_html_json = True
                html_content = parsed["data"]
            else:
                content = json.dumps(parsed, indent=4)
        except Exception:
            pass
        if is_html_json:
            doc.setHtml(html_content)
        else:
            escaped = html.escape(content)
            doc.setHtml(
                f'<html><body style="font-family: Consolas, \'Courier New\', monospace; font-size: 9pt; white-space: pre-wrap;">{escaped}</body></html>'
            )
    elif ext in (".html", ".htm"):
        doc.setHtml(content)
    elif ext in (".md", ".markdown"):
        doc.setMarkdown(content)
    else:
        escaped = html.escape(content)
        doc.setHtml(
            f'<html><body style="font-family: Consolas, \'Courier New\', monospace; font-size: 9.5pt; white-space: pre-wrap; line-height: 1.4;">{escaped}</body></html>'
        )

    try:
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(pdf_path)

        layout = QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4),
            QPageLayout.Orientation.Portrait,
            QMarginsF(15, 15, 15, 15),
            QPageLayout.Unit.Millimeter,
        )
        printer.setPageLayout(layout)
        doc.print_(printer)
        if os.path.isfile(pdf_path) and os.path.getsize(pdf_path) > 0:
            return
    except Exception as q_err:
        _log.warning("QTextDocument print failed for %s: %s; falling back to PyMuPDF.", text_path, q_err)

    pdf_doc = fitz.open()
    try:
        process_txt(pdf_doc, text_path)
        if pdf_doc.page_count > 0:
            atomic_pdf_write(pdf_doc, pdf_path, close_writer=True)
            return
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_doc, "is_closed", False):
                pdf_doc.close()

    raise RuntimeError(f"Could not convert text document: {text_path}")


# ── Internal Document Stream Processors ──────────────────────────────────────

def process_image(doc: fitz.Document, img_path: str) -> bool:
    """Insert an image as a new page matching its exact aspect ratio."""
    ok, _w, _h = check_image_size(img_path)
    if not ok:
        return False
    try:
        img = fitz.open(img_path)
        try:
            if img.page_count == 0:
                return False
            rect = img[0].rect
            page = doc.new_page(width=rect.width, height=rect.height)
            page.insert_image(page.rect, filename=img_path)
            return True
        finally:
            with contextlib.suppress(Exception):
                if not getattr(img, "is_closed", False):
                    img.close()
    except Exception:
        try:
            with Image.open(img_path) as pimg:
                pimg = ImageOps.exif_transpose(pimg)
                w, h = pimg.size
            page = doc.new_page(width=w, height=h)
            page.insert_image(page.rect, filename=img_path)
            return True
        except Exception:
            return False


def process_txt(doc: fitz.Document, txt_path: str) -> None:
    """Convert text files into multipage PDF documents with line wrapping."""
    with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    lines = content.split("\n")
    page = None
    y = 50.0
    fontsize = 10.0
    line_height = fontsize * 1.4
    margin_x = 50.0
    max_y = 792.0
    max_width = 495.0
    for line in lines:
        if page is None or y + fontsize > max_y:
            page = doc.new_page(width=595, height=842)
            y = 50.0
        if not line.strip():
            y += line_height
            continue
        rect = fitz.Rect(margin_x, y, margin_x + max_width, max_y)
        used = page.insert_textbox(rect, line, fontsize=fontsize, fontname="helv")
        if used < 0:
            page = doc.new_page(width=595, height=842)
            y = 50.0
            rect = fitz.Rect(margin_x, y, margin_x + max_width, max_y)
            used = page.insert_textbox(rect, line, fontsize=fontsize, fontname="helv")
        est_lines = max(1, len(line) * fontsize * 0.5 / max_width + 1)
        y += line_height * est_lines


def md_to_lines(md: str) -> list[tuple[str, int, bool]]:
    result = []
    for line in md.split("\n"):
        stripped = line.strip()
        if stripped.startswith("# "):
            result.append((stripped[2:], 18, True))
        elif stripped.startswith("## "):
            result.append((stripped[3:], 15, True))
        elif stripped.startswith("### "):
            result.append((stripped[4:], 13, True))
        elif stripped.startswith("#### "):
            result.append((stripped[5:], 12, True))
        elif stripped.startswith("- ") or stripped.startswith("* "):
            result.append(("  \u2022  " + stripped[2:], 10, False))
        elif stripped.startswith("```"):
            continue
        elif stripped in ("---", "***"):
            result.append(("\u2500" * 60, 8, False))
        elif stripped == "":
            result.append(("", 10, False))
        else:
            clean = stripped.replace("**", "").replace("__", "")
            clean = clean.replace("*", "").replace("_", "")
            clean = clean.replace("`", "")
            result.append((clean, 10, False))
    return result


def html_to_lines(soup: Any) -> list[tuple[str, int, bool]]:
    lines = []
    tag_map = {
        "h1": (18, True),
        "h2": (15, True),
        "h3": (13, True),
        "h4": (12, True),
        "h5": (11, True),
        "h6": (11, True),
    }
    for el in soup.find_all(
        [
            "h1", "h2", "h3", "h4", "h5", "h6",
            "p", "li", "pre", "td", "th", "blockquote",
        ]
    ):
        text = el.get_text(strip=True)
        if not text:
            continue
        tag = el.name
        if tag in tag_map:
            size, bold = tag_map[tag]
            lines.append((text, size, bold))
        elif tag == "li":
            lines.append(("  \u2022  " + text, 10, False))
        elif tag == "pre":
            for ln in text.split("\n"):
                lines.append((ln, 9, False))
        elif tag == "blockquote":
            lines.append(("  \u201c " + text + " \u201d", 10, False))
        elif tag in ("th",):
            lines.append((text, 10, True))
        else:
            lines.append((text, 10, False))
        if tag in tag_map:
            lines.append(("", 10, False))
    return lines


def docx_xml_to_lines(docx_path: str) -> list[tuple[str, int, bool]]:
    """Native fallback to parse word/document.xml directly without python-docx."""
    lines: list[tuple[str, int, bool]] = []
    try:
        with zipfile.ZipFile(docx_path) as z:
            if "word/document.xml" not in z.namelist():
                return lines
            xml_bytes = z.read("word/document.xml")
        root = ET.fromstring(xml_bytes)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for p in root.iter(f"{{{ns['w']}}}p"):
            texts = [t.text for t in p.iter(f"{{{ns['w']}}}t") if t.text]
            p_text = "".join(texts).strip()
            if p_text:
                is_bold = any(p.iter(f"{{{ns['w']}}}b"))
                p_style = ""
                for s in p.iter(f"{{{ns['w']}}}pStyle"):
                    val = s.attrib.get(f"{{{ns['w']}}}val", "")
                    if val:
                        p_style = val.lower()
                size = 10
                if "heading1" in p_style or "heading 1" in p_style:
                    size = 18
                    is_bold = True
                elif "heading2" in p_style or "heading 2" in p_style:
                    size = 15
                    is_bold = True
                elif "heading3" in p_style or "heading 3" in p_style:
                    size = 13
                    is_bold = True
                elif "heading" in p_style:
                    size = 12
                    is_bold = True
                lines.append((p_text, size, is_bold))
            else:
                lines.append(("", 10, False))
    except Exception as exc:
        _log.warning("docx_xml_to_lines fallback error: %s", exc)
    return lines


def pptx_xml_to_lines(pptx_path: str) -> list[tuple[str, int, bool]]:
    """Native fallback to parse ppt/slides/slide*.xml directly without python-pptx."""
    lines: list[tuple[str, int, bool]] = []
    try:
        with zipfile.ZipFile(pptx_path) as z:
            slide_files = sorted(
                [name for name in z.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
            )
            for s_file in slide_files:
                xml_bytes = z.read(s_file)
                root = ET.fromstring(xml_bytes)
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                for p in root.iter(f"{{{ns['a']}}}p"):
                    texts = [t.text for t in p.iter(f"{{{ns['a']}}}t") if t.text]
                    p_text = "".join(texts).strip()
                    if p_text:
                        lines.append((p_text, 12, False))
                lines.append(("", 10, False))
    except Exception as exc:
        _log.warning("pptx_xml_to_lines fallback error: %s", exc)
    return lines


def xlsx_xml_to_lines(xlsx_path: str) -> list[tuple[str, int, bool]]:
    """Native fallback to parse xl/worksheets/sheet*.xml directly without openpyxl."""
    lines: list[tuple[str, int, bool]] = []
    try:
        with zipfile.ZipFile(xlsx_path) as z:
            shared_strings: list[str] = []
            if "xl/sharedStrings.xml" in z.namelist():
                ss_bytes = z.read("xl/sharedStrings.xml")
                ss_root = ET.fromstring(ss_bytes)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for si in ss_root.iter(f"{{{ns['s']}}}si"):
                    texts = [t.text for t in si.iter(f"{{{ns['s']}}}t") if t.text]
                    shared_strings.append("".join(texts))

            sheet_files = sorted(
                [name for name in z.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")]
            )
            ns_s = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            for sf in sheet_files:
                xml_bytes = z.read(sf)
                root = ET.fromstring(xml_bytes)
                for row in root.iter(f"{{{ns_s['s']}}}row"):
                    row_vals = []
                    for c in row.iter(f"{{{ns_s['s']}}}c"):
                        cell_type = c.attrib.get("t", "")
                        v_el = c.find(f"{{{ns_s['s']}}}v")
                        val = v_el.text if v_el is not None and v_el.text else ""
                        if cell_type == "s" and val.isdigit():
                            idx = int(val)
                            val = shared_strings[idx] if idx < len(shared_strings) else val
                        row_vals.append(val)
                    if any(row_vals):
                        lines.append(("  |  ".join(row_vals), 9, False))
                lines.append(("", 10, False))
    except Exception as exc:
        _log.warning("xlsx_xml_to_lines fallback error: %s", exc)
    return lines


def render_lines_to_doc(doc: fitz.Document, lines: list[tuple[str, int, bool]]) -> None:
    """Render paragraph tuples (text, size, bold) into multi-page PDF documents with word wrapping."""
    page = None
    y = 50.0
    margin_x = 50.0
    max_w = 495.0
    max_y = 792.0

    for text, size, bold in lines:
        if not text:
            y += size * 0.8
            continue

        font = "hebo" if bold else "helv"
        line_height = size * 1.35

        if page is None or y + line_height > max_y:
            page = doc.new_page(width=595, height=842)
            y = 50.0

        rect = fitz.Rect(margin_x, y, margin_x + max_w, max_y)
        rc = page.insert_textbox(rect, text, fontsize=size, fontname=font)
        if rc < 0:
            page = doc.new_page(width=595, height=842)
            y = 50.0
            rect = fitz.Rect(margin_x, y, margin_x + max_w, max_y)
            page.insert_textbox(rect, text, fontsize=size, fontname=font)

        est_chars_per_line = max(10, int(max_w / (size * 0.52)))
        est_lines = max(1, (len(text) + est_chars_per_line - 1) // est_chars_per_line)
        y += est_lines * line_height + 4.0


def process_md(doc: fitz.Document, md_path: str) -> None:
    with open(md_path, "r", encoding="utf-8", errors="replace") as f:
        md_text = f.read()
    lines = md_to_lines(md_text)
    render_lines_to_doc(doc, lines)


def process_docx(doc: fitz.Document, docx_path: str) -> None:
    lines: list[tuple[str, int, bool]] = []
    try:
        from docx import Document
        dx = Document(docx_path)
        for para in dx.paragraphs:
            text = para.text.strip()
            p_style = para.style
            style = (p_style.name or "").lower() if p_style is not None else ""
            if "heading 1" in style:
                lines.append((text, 18, True))
            elif "heading 2" in style:
                lines.append((text, 15, True))
            elif "heading 3" in style:
                lines.append((text, 13, True))
            elif "heading" in style:
                lines.append((text, 12, True))
            elif text:
                bold = any(r.bold for r in para.runs if r.bold)
                lines.append((text, 10, bold))
            else:
                lines.append(("", 10, False))
        for table in dx.tables:
            lines.append(("", 10, False))
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(("  |  ".join(cells), 9, False))
            lines.append(("", 10, False))
    except Exception as exc:
        _log.warning("python-docx processing failed for %s: %s; using XML fallback.", docx_path, exc)
        lines = []

    if not lines or not any(txt for txt, _, _ in lines):
        fallback_lines = docx_xml_to_lines(docx_path)
        if fallback_lines:
            lines = fallback_lines

    if lines and any(txt for txt, _, _ in lines):
        render_lines_to_doc(doc, lines)

    # If still 0 pages (e.g. document only has images or shapes), extract images from word/media/
    if doc.page_count == 0:
        try:
            with zipfile.ZipFile(docx_path) as z:
                media_files = sorted([m for m in z.namelist() if m.startswith("word/media/")])
                for mf in media_files:
                    img_bytes = z.read(mf)
                    try:
                        ext = os.path.splitext(mf)[1].lower().lstrip(".")
                        img = fitz.open(stream=img_bytes, filetype=ext)
                        if img.page_count > 0:
                            r = img[0].rect
                            page = doc.new_page(width=r.width, height=r.height)
                            page.insert_image(page.rect, stream=img_bytes)
                        img.close()
                    except Exception:
                        pass
        except Exception:
            pass

    if doc.page_count == 0:
        doc.new_page(width=595, height=842)


def process_pptx(doc: fitz.Document, pptx_path: str) -> None:
    try:
        from pptx import Presentation
        prs = Presentation(pptx_path)
        slide_w = prs.slide_width.pt if prs.slide_width else 960
        slide_h = prs.slide_height.pt if prs.slide_height else 540
        for slide in prs.slides:
            page = doc.new_page(width=slide_w, height=slide_h)
            y = 40.0
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                tf = getattr(shape, "text_frame", None)
                if tf is None:
                    continue
                for para in tf.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    size = 10.0
                    bold = False
                    if para.runs:
                        r = para.runs[0]
                        font = getattr(r, "font", None)
                        if font is not None and getattr(font, "size", None):
                            size = min(36.0, float(font.size.pt))
                        if font is not None:
                            bold = bool(getattr(font, "bold", False))
                    font_name = "hebo" if bold else "helv"
                    if y + size > slide_h - 20:
                        break
                    try:
                        page.insert_text(fitz.Point(40, y), text, fontsize=size, fontname=font_name)
                    except Exception:
                        page.insert_text(fitz.Point(40, y), text, fontsize=size, fontname="helv")
                    y += size * 1.4
    except Exception:
        fallback_lines = pptx_xml_to_lines(pptx_path)
        if fallback_lines:
            render_lines_to_doc(doc, fallback_lines)
    if doc.page_count == 0:
        doc.new_page(width=960, height=540)


def process_xlsx(doc: fitz.Document, xlsx_path: str) -> None:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, data_only=True)
        for ws in wb.worksheets:
            lines = []
            lines.append((f"— {ws.title} —", 12, True))
            lines.append(("", 10, False))
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(("  |  ".join(cells), 8, False))
            lines.append(("", 10, False))
            render_lines_to_doc(doc, lines)
    except Exception:
        fallback_lines = xlsx_xml_to_lines(xlsx_path)
        if fallback_lines:
            render_lines_to_doc(doc, fallback_lines)
    if doc.page_count == 0:
        doc.new_page(width=595, height=842)


def process_html(doc: fitz.Document, html_path: str) -> None:
    try:
        from bs4 import BeautifulSoup
        with open(html_path, "r", encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        lines = html_to_lines(soup)
        render_lines_to_doc(doc, lines)
    except Exception:
        fallback_lines = []
        with open(html_path, "r", encoding="utf-8", errors="replace") as f:
            raw = f.read()
        clean = re.sub(r"<[^>]+>", " ", raw).strip()
        for line in clean.split("\n"):
            s = line.strip()
            if s:
                fallback_lines.append((s, 10, False))
        if fallback_lines:
            render_lines_to_doc(doc, fallback_lines)
        else:
            raise


def process_epub(doc: fitz.Document, epub_path: str) -> None:
    converted = False
    try:
        epub_doc = fitz.open(epub_path)
        try:
            pdf_bytes = epub_doc.convert_to_pdf()
            pdf_part = fitz.open("pdf", pdf_bytes)
            try:
                doc.insert_pdf(pdf_part)
            finally:
                with contextlib.suppress(Exception):
                    if not getattr(pdf_part, "is_closed", False):
                        pdf_part.close()
            converted = True
        finally:
            with contextlib.suppress(Exception):
                if not getattr(epub_doc, "is_closed", False):
                    epub_doc.close()
    except Exception:
        converted = False
    if not converted:
        try:
            epub_doc = fitz.open(epub_path)
            try:
                lines = []
                for page in epub_doc:
                    raw_text = page.get_text("text")
                    text = str(raw_text or "")
                    for line in text.split("\n"):
                        if line.strip():
                            lines.append((line.strip(), 10, False))
                        else:
                            lines.append(("", 10, False))
                if lines:
                    render_lines_to_doc(doc, lines)
            finally:
                with contextlib.suppress(Exception):
                    if not getattr(epub_doc, "is_closed", False):
                        epub_doc.close()
        except Exception:
            pass


def process_pdf(doc: fitz.Document, pdf_path: str) -> None:
    pdf_part = fitz.open(pdf_path)
    try:
        doc.insert_pdf(pdf_part)
    finally:
        with contextlib.suppress(Exception):
            if not getattr(pdf_part, "is_closed", False):
                pdf_part.close()


# ── Unified Converter Facade ───────────────────────────────────────────────────

def convert_file_to_pdf(file_path: str, out_pdf_path: str | None = None) -> str:
    """Convert any supported document, image, or text file into a PDF without user intervention."""
    if not file_path or not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    is_pdf = file_path.lower().endswith(".pdf")
    if not is_pdf:
        try:
            with open(file_path, "rb") as f:
                header = f.read(1024)
                if b"%PDF-" in header:
                    return file_path
        except Exception:
            pass
    else:
        return file_path

    if out_pdf_path is None:
        fd, out_pdf_path = tempfile.mkstemp(prefix="pdfapps_conv_", suffix=".pdf")
        os.close(fd)

    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext in DOC_EXTS:
            convert_docx_to_pdf(file_path, out_pdf_path)
        elif ext in PPT_EXTS:
            convert_pptx_to_pdf(file_path, out_pdf_path)
        elif ext in EXCEL_EXTS:
            convert_xlsx_to_pdf(file_path, out_pdf_path)
        elif ext in IMAGE_EXTS:
            convert_image_to_pdf(file_path, out_pdf_path)
        elif ext in EBOOK_EXTS:
            converted = False
            try:
                src_doc = fitz.open(file_path)
                try:
                    if hasattr(src_doc, "convert_to_pdf"):
                        pdf_bytes = src_doc.convert_to_pdf()
                        pdf_part = fitz.open("pdf", pdf_bytes)
                        try:
                            atomic_pdf_write(pdf_part, out_pdf_path, close_writer=True)
                            converted = True
                        finally:
                            with contextlib.suppress(Exception):
                                if not getattr(pdf_part, "is_closed", False):
                                    pdf_part.close()
                finally:
                    with contextlib.suppress(Exception):
                        if not getattr(src_doc, "is_closed", False):
                            src_doc.close()
            except Exception as e_conv:
                _log.warning("PyMuPDF convert_to_pdf failed for %s: %s", file_path, e_conv)
            if not converted:
                convert_text_to_pdf(file_path, out_pdf_path)
        elif ext in TEXT_EXTS:
            convert_text_to_pdf(file_path, out_pdf_path)
        else:
            try:
                convert_text_to_pdf(file_path, out_pdf_path)
            except Exception:
                doc = fitz.open()
                try:
                    process_txt(doc, file_path)
                    atomic_pdf_write(doc, out_pdf_path, close_writer=True)
                finally:
                    with contextlib.suppress(Exception):
                        if not getattr(doc, "is_closed", False):
                            doc.close()

        if not os.path.isfile(out_pdf_path) or os.path.getsize(out_pdf_path) == 0:
            raise ValueError(f"No pages could be generated from '{file_path}'")

    except Exception:
        if os.path.exists(out_pdf_path):
            with contextlib.suppress(Exception):
                os.unlink(out_pdf_path)
        raise

    return out_pdf_path


# ── TabImport GUI Component ──────────────────────────────────────────────────

class _ImportFileList(QListWidget):
    """File list supporting drag & drop of arbitrary files directly into the list."""

    def __init__(self, on_drop_callback, parent=None):
        super().__init__(parent)
        self._on_drop_callback = on_drop_callback
        self.setAcceptDrops(True)
        self.setDragDropMode(QAbstractItemView.DragDropMode.InternalMove)
        self.setAlternatingRowColors(True)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragEnterEvent(event)

    def dragMoveEvent(self, event):
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
        else:
            super().dragMoveEvent(event)

    def dropEvent(self, event):
        if event.mimeData().hasUrls():
            paths = [
                os.path.abspath(u.toLocalFile())
                for u in event.mimeData().urls()
                if u.toLocalFile() and os.path.isfile(u.toLocalFile())
            ]
            if paths:
                self._on_drop_callback(paths)
                event.acceptProposedAction()
                return
        super().dropEvent(event)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Delete, Qt.Key.Key_Backspace):
            for item in self.selectedItems():
                row = self.row(item)
                self.takeItem(row)
            event.accept()
            return
        super().keyPressEvent(event)


class TabImport(BasePage):
    """Import and convert any supported document or image types directly to PDF with auto-detection."""

    def __init__(self, status_fn):
        super().__init__(
            "fa5s.file-import",
            t("tool.import.name"),
            t("tool.import.desc"),
            t("tool.import.btn"),
            status_fn,
        )
        f = self._form

        f.addWidget(section(t("tool.import.source_file")))

        self._file_list = _ImportFileList(self._on_files_dropped, self)
        self._file_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self._file_list.setMinimumHeight(180)
        f.addWidget(self._file_list)

        file_btns = QHBoxLayout()
        self._add_btn = QPushButton(t("tool.import.add_files"))
        self._add_btn.setIcon(qta.icon("fa5s.plus", color=TEXT_SEC))
        self._add_btn.clicked.connect(self._pick_files)
        file_btns.addWidget(self._add_btn)

        self._remove_selected_btn = QPushButton(t("btn.remove"))
        self._remove_selected_btn.setIcon(qta.icon("fa5s.minus", color=TEXT_SEC))
        self._remove_selected_btn.clicked.connect(self._remove_selected_files)
        file_btns.addWidget(self._remove_selected_btn)

        self._clear_btn = danger_btn(t("btn.clear"))
        self._clear_btn.clicked.connect(self._file_list.clear)
        file_btns.addWidget(self._clear_btn)
        file_btns.addStretch()
        f.addLayout(file_btns)

        f.addWidget(section(t("tool.import.output")))
        self.drop_out = DropFileEdit(save=True, default_name="output.pdf")
        f.addWidget(self.drop_out)

        self.lbl_result = QLabel("")
        self.lbl_result.setStyleSheet(result_label_style())
        f.addWidget(self.lbl_result)
        f.addStretch()

    def set_compact_mode(self, active: bool, path: str = "") -> None:
        pass

    def update_theme(self, dark: bool) -> None:
        super().update_theme(dark)
        try:
            self.lbl_result.setStyleSheet(result_label_style(dark))
        except RuntimeError:
            pass

    @staticmethod
    def _icon_for_ext(ext: str) -> str:
        ext = ext.lower()
        if ext in IMAGE_EXTS:
            return "fa5s.image"
        if ext in DOC_EXTS:
            return "fa5s.file-word"
        if ext in PPT_EXTS:
            return "fa5s.file-powerpoint"
        if ext in EXCEL_EXTS:
            return "fa5s.file-excel"
        if ext in (".html", ".htm"):
            return "fa5s.file-code"
        if ext in EBOOK_EXTS:
            return "fa5s.book"
        if ext in (".pdf",):
            return "fa5s.file-pdf"
        return "fa5s.file-alt"

    def _add_file_path(self, p: str):
        if not p or not os.path.isfile(p):
            return
        p = os.path.abspath(p)
        existing = [
            self._file_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self._file_list.count())
        ]
        if p in existing:
            return

        ext = os.path.splitext(p)[1].lower()
        icon_name = self._icon_for_ext(ext)
        item = QListWidgetItem(qta.icon(icon_name, color=TEXT_SEC), os.path.basename(p))
        item.setData(Qt.ItemDataRole.UserRole, p)
        item.setToolTip(p)
        self._file_list.addItem(item)

        if not self.drop_out.path():
            base = os.path.splitext(p)[0]
            self.drop_out.set_path(base + ".pdf")

    def _on_files_dropped(self, paths: list[str]):
        for p in paths:
            self._add_file_path(p)

    def _pick_files(self):
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            t("tool.import.add_files"),
            DESKTOP,
            _ALL_SUPPORTED_FILTER,
        )
        for p in paths:
            self._add_file_path(p)

    def _remove_selected_files(self):
        for item in self._file_list.selectedItems():
            row = self._file_list.row(item)
            self._file_list.takeItem(row)

    def _get_files(self) -> list[str]:
        return [
            self._file_list.item(i).data(Qt.ItemDataRole.UserRole)
            or self._file_list.item(i).text()
            for i in range(self._file_list.count())
        ]

    def _run(self):
        files = self._get_files()
        if not files:
            QMessageBox.warning(self, t("msg.warning"), t("tool.import.select_file"))
            return

        default_name = "output.pdf"
        if files:
            base, _ = os.path.splitext(os.path.basename(files[0]))
            default_name = f"{base}.pdf"
        start_dir = os.path.dirname(files[0]) if files else ""
        out = self._prompt_save_as(default_name, start_dir)
        if not out:
            return
        self.drop_out.set_path(out)

        self.lbl_result.setText("")
        n = len(files)

        def do_work(worker):
            doc = fitz.open()
            skipped = 0
            try:
                for i, file_path in enumerate(files):
                    if worker.is_cancelled():
                        return None
                    worker.progress.emit(i + 1, f"{i + 1}/{n}: {os.path.basename(file_path)}")
                    try:
                        part_pdf = convert_file_to_pdf(file_path)
                        part_doc = fitz.open(part_pdf)
                        doc.insert_pdf(part_doc)
                        with contextlib.suppress(Exception):
                            if not getattr(part_doc, "is_closed", False):
                                part_doc.close()
                        if part_pdf != file_path and os.path.isfile(part_pdf):
                            with contextlib.suppress(Exception):
                                os.unlink(part_pdf)
                    except Exception as exc_part:
                        _log.warning("Failed to convert %s: %s", file_path, exc_part)
                        skipped += 1

                if worker.is_cancelled():
                    return None
                if doc.page_count == 0:
                    return _NoContent(skipped)

                win: Any = self.window()
                viewer = getattr(win, "_viewer", None)
                if (
                    viewer
                    and viewer.current_path()
                    and os.path.abspath(viewer.current_path()) == os.path.abspath(out)
                ):
                    viewer._canvas.close_doc()
                    if viewer._fitz_doc:
                        with contextlib.suppress(Exception):
                            viewer._fitz_doc.close()
                        viewer._fitz_doc = None
                    viewer._thumbnails._stop_all_workers()

                atomic_pdf_write(doc, out, close_writer=True)
            finally:
                with contextlib.suppress(Exception):
                    if not getattr(doc, "is_closed", False):
                        doc.close()
            return out

        self._run_background(
            do_work,
            total=max(n, 1),
            label=t("tool.import.converting"),
            on_done=lambda r: self._on_result(r, out),
        )

    def _on_result(self, result, out_path: str):
        if isinstance(result, _NoContent):
            if result.skipped:
                self._status(t("tool.import.skipped_images", n=result.skipped))
            else:
                self._status(t("tool.import.no_content"))
            QMessageBox.warning(self, t("msg.warning"), t("tool.import.no_content"))
            return
        self._done(out_path)

    def _done(self, out_path: str):
        self.lbl_result.setText(f"  \u2192 {os.path.basename(out_path)}")
        self._status(t("tool.import.status.done", path=out_path))

        win: Any = self.window()
        viewer = getattr(win, "_viewer", None)
        cleanup_fn = getattr(win, "_cleanup_pipeline", None)
        if callable(cleanup_fn) and viewer:
            cleanup_fn(id(viewer))

        if viewer:
            viewer.load(out_path)

        QMessageBox.information(self, t("msg.done"), t("tool.import.done", path=out_path))


__all__ = [
    "TabImport",
    "convert_file_to_pdf",
    "convert_docx_to_pdf",
    "convert_pptx_to_pdf",
    "convert_xlsx_to_pdf",
    "convert_image_to_pdf",
    "convert_text_to_pdf",
    "process_image",
    "process_txt",
    "process_md",
    "process_docx",
    "process_pptx",
    "process_xlsx",
    "process_html",
    "process_epub",
    "process_pdf",
]